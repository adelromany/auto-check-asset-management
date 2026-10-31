"""Tests for the Print Laptop Form feature (Windows printing).

Automated tests mock the Windows print system at the service seam
(`app.services.printing.run_powershell` / `run_ghostscript`) and the heavy
pipeline (`assemble_and_print`), per docs/windows-printing.md ("Test document
generation and printer submission separately. Use a mocked printing service
for automated tests"). The real print test is manual on the server.
"""

import json

import docx
import pytest

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _fake_powershell_json(printers):
    """Return a run_powershell stand-in emitting Get-Printer JSON output."""
    def _run(script, timeout=None):
        assert "Get-Printer" in script
        return json.dumps(printers)
    return _run


SAMPLE_DATA = {
    "hostname": "LT-0042",
    "username": "j.doe",
    "serial": "SN-XYZ-9911",
    "manufacturer": "Dell Inc.",
    "model": "Latitude 5440",
    "cpu": "Intel i7-1365U",
    "cpu_cores": 10,
    "ram_gb": 16,
    "disks": [{"name": "Samsung SSD", "size_gb": 512, "type": "NVMe"}],
    "os_name": "Windows 11 Pro",
    "os_version": "10.0.26100",
    "os_arch": "64-bit",
    "ip_addresses": [{"iface": "Ethernet", "ip": "10.42.1.7"}],
    "agent_os": "windows",
    "collected_at": "2026-10-20T10:00:00+00:00",
    "software_checks": [
        {"id": "chrome", "name": "Google Chrome", "installed": True,
         "version": "131.0.6778.85", "method": "registry"},
    ],
    "windows_activation": {"status": "Licensed", "product_name": "Windows 11 Pro",
                           "partial_key": "XYZ12"},
}


@pytest.fixture()
def session_with_data(sso_off_client, monkeypatch):
    """A fresh session with agent data already submitted."""
    client = sso_off_client
    code = client.get("/api/session/new").json()["code"]
    client.post(f"/api/session/{code}/submit", json=SAMPLE_DATA)
    return client, code


# ---------------------------------------------------------------------------
# GET /api/print/printers
# ---------------------------------------------------------------------------

def test_printers_endpoint_mocked(make_client, monkeypatch):
    client, _ = make_client()
    import app.services.printing as printing
    monkeypatch.setattr(
        printing, "run_powershell",
        _fake_powershell_json([{"Name": "HP LaserJet", "PrinterStatus": 0,
                                "WorkOffline": False, "Type": 0}]))
    resp = client.get("/api/print/printers")
    assert resp.status_code == 200
    body = resp.json()
    assert body["printers"] == [{"name": "HP LaserJet", "type": 0,
                                 "status": 0, "work_offline": False}]
    assert body["printer_names"] == ["HP LaserJet"]
    assert body["configured"] == ""           # nothing admin-configured yet
    assert body["selector_enabled"] is False
    assert body["template_found"] is True     # starter template ships in repo


def test_printers_endpoint_single_queue_json_dict(make_client, monkeypatch):
    """PowerShell's ConvertTo-Json unwraps 1-element lists — must be handled."""
    client, _ = make_client()
    import app.services.printing as printing
    monkeypatch.setattr(
        printing, "run_powershell",
        _fake_powershell_json({"Name": "Lone Printer", "PrinterStatus": 0,
                               "WorkOffline": True, "Type": 0}))
    body = client.get("/api/print/printers").json()
    assert body["printers"][0]["name"] == "Lone Printer"
    assert body["printers"][0]["work_offline"] is True


def test_printers_endpoint_503_non_windows(make_client, monkeypatch):
    client, _ = make_client()
    import app.services.printing as printing

    def _boom(script, timeout=None):
        raise printing.PrintingError(
            "powershell.exe not found. Printing requires Windows Server.")
    monkeypatch.setattr(printing, "run_powershell", _boom)
    resp = client.get("/api/print/printers")
    assert resp.status_code == 503
    assert "Windows" in resp.json()["detail"]


def test_print_endpoints_401_when_sso_on(make_client):
    client, _ = make_client(SSO_ENABLED="true")
    assert client.get("/api/print/printers").status_code == 401
    assert client.get("/api/print/jobs?printer=x").status_code == 401
    assert client.post("/api/session/ABC/print").status_code == 401


# ---------------------------------------------------------------------------
# POST /api/session/{code}/print — guard rails
# ---------------------------------------------------------------------------

def test_print_unknown_session_404(sso_off_client):
    resp = sso_off_client.post("/api/session/NOPE/print")
    assert resp.status_code == 404


def test_print_without_agent_data_409(sso_off_client):
    client = sso_off_client
    code = client.get("/api/session/new").json()["code"]
    resp = client.post(f"/api/session/{code}/print")
    assert resp.status_code == 409
    assert "run the agent" in resp.json()["detail"]


def test_print_no_printer_configured_503(session_with_data, monkeypatch):
    client, code = session_with_data
    resp = client.post(f"/api/session/{code}/print")
    assert resp.status_code == 503
    assert "PRINT_PRINTER" in resp.json()["detail"]


def test_print_selector_disabled_rejects_body_printer(session_with_data):
    client, code = session_with_data
    resp = client.post(f"/api/session/{code}/print", json={"printer": "HP"})
    assert resp.status_code == 403
    assert "disabled" in resp.json()["detail"]


def test_print_selector_enabled_rejects_unknown_printer(make_client, monkeypatch):
    client, mod = make_client(PRINT_PRINTER="HP LaserJet",
                              PRINT_ALLOW_PRINTER_SELECTOR="true")
    code = client.get("/api/session/new").json()["code"]
    client.post(f"/api/session/{code}/submit", json=SAMPLE_DATA)

    import app.services.printing as printing
    monkeypatch.setattr(
        printing, "run_powershell",
        _fake_powershell_json([{"Name": "Other Printer", "PrinterStatus": 0,
                                "WorkOffline": False, "Type": 0}]))
    resp = client.post(f"/api/session/{code}/print", json={"printer": "Phantom"})
    assert resp.status_code == 503
    assert "not installed" in resp.json()["detail"]


def test_print_template_missing_503(session_with_data, monkeypatch):
    client, code = session_with_data
    import app.config as config
    monkeypatch.setattr(config, "PRINT_PRINTER", "HP LaserJet")
    absent = "nonexistent-template.docx"
    monkeypatch.setattr(config, "FORM_TEMPLATE_PATH", absent)
    resp = client.post(f"/api/session/{code}/print")
    assert resp.status_code == 503
    assert "template" in resp.json()["detail"].lower()
    # actionable, without leaking absolute server filesystem paths
    assert absent not in resp.json()["detail"]


# ---------------------------------------------------------------------------
# POST /api/session/{code}/print — happy path (pipeline mocked)
# ---------------------------------------------------------------------------

def test_print_happy_path(session_with_data, monkeypatch):
    client, code = session_with_data
    import app.config as config
    import app.services.laptop_form as laptop_form
    import app.services.printing as printing

    monkeypatch.setattr(config, "PRINT_PRINTER", "HP LaserJet")

    calls = {}

    def fake_assemble(data, snipe, printer_name, session_code):
        calls.update({"data": data, "snipe": snipe, "printer": printer_name,
                      "code": session_code})
        return {
            "status": "submitted", "printer": printer_name,
            "document_name": "AssetForm_X", "docx": "AssetForm_X.docx",
            "pdf": "AssetForm_X.pdf", "submitted_at": "2026-10-20T00:00:00+00:00",
            "job": {"job_id": 3, "status": 0}, "tokens_missing": [],
            "note": "Accepted by print spooler",
        }

    monkeypatch.setattr(laptop_form, "assemble_and_print", fake_assemble)

    resp = client.post(f"/api/session/{code}/print")
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "submitted"
    assert body["printer"] == "HP LaserJet"
    assert body["job"] == {"job_id": 3, "status": 0}
    # the pipeline received the session's latest inventory verbatim
    assert calls["data"]["serial"] == "SN-XYZ-9911"
    assert calls["printer"] == "HP LaserJet"
    assert calls["code"] == code
    # no Snipe-IT token configured → enrichment is best-effort None, still ok
    assert calls["snipe"] is None
    # response must not leak server paths or commands
    flat = json.dumps(body).lower()
    assert "python" not in flat and "\\\\" not in flat and "powershell" not in flat


def test_print_pipeline_error_maps_503(session_with_data, monkeypatch):
    client, code = session_with_data
    import app.config as config
    import app.services.laptop_form as laptop_form

    monkeypatch.setattr(config, "PRINT_PRINTER", "HP LaserJet")

    def boom(data, snipe, printer, code):
        raise laptop_form.LaptopFormError(
            "LibreOffice could not convert the document to PDF.")
    monkeypatch.setattr(laptop_form, "assemble_and_print", boom)
    resp = client.post(f"/api/session/{code}/print")
    assert resp.status_code == 503
    assert "LibreOffice" in resp.json()["detail"]


def test_print_spooler_error_maps_502(session_with_data, monkeypatch):
    client, code = session_with_data
    import app.config as config
    import app.services.laptop_form as laptop_form
    import app.services.printing as printing

    monkeypatch.setattr(config, "PRINT_PRINTER", "HP LaserJet")

    def boom(data, snipe, printer, code):
        raise printing.PrintingError("Ghostscript exited with code 1.")
    monkeypatch.setattr(laptop_form, "assemble_and_print", boom)
    resp = client.post(f"/api/session/{code}/print")
    assert resp.status_code == 502
    assert "Ghostscript" in resp.json()["detail"]


# ---------------------------------------------------------------------------
# GET /api/print/jobs
# ---------------------------------------------------------------------------

def test_jobs_endpoint_mocked(make_client, monkeypatch):
    client, _ = make_client()
    import app.services.printing as printing
    monkeypatch.setattr(
        printing, "run_powershell",
        _fake_powershell_json([{"Id": 7, "DocumentName": "AssetForm_X.pdf",
                                "JobStatus": 64, "UserName": "svc-snipeit",
                                "SubmittedTime": "2026-10-20T10:01:02"}]))
    resp = client.get("/api/print/jobs", params={"printer": "HP LaserJet"})
    assert resp.status_code == 200
    job = resp.json()["jobs"][0]
    assert job["job_id"] == 7
    assert job["job_status"] == 64
    assert job["document_name"] == "AssetForm_X.pdf"


def test_jobs_endpoint_requires_printer(make_client):
    client, _ = make_client()
    resp = client.get("/api/print/jobs")
    assert resp.status_code == 422       # FastAPI validation rejects missing
                                         # query param before our handler


def test_jobs_endpoint_blank_printer_400(make_client):
    client, _ = make_client()
    resp = client.get("/api/print/jobs", params={"printer": ""})
    assert resp.status_code == 400


def test_jobs_endpoint_unknown_printer_503(make_client, monkeypatch):
    client, _ = make_client()
    import app.services.printing as printing

    def _run(script, timeout=None):
        raise printing.PrintingError(
            "Windows print system error: Printer does not exist")
    monkeypatch.setattr(printing, "run_powershell", _run)
    resp = client.get("/api/print/jobs", params={"printer": "Ghost"})
    assert resp.status_code == 503


# ---------------------------------------------------------------------------
# document generation unit tests (real python-docx, no Windows printing)
# ---------------------------------------------------------------------------

def test_build_mapping_formats_lists():
    from app.services.laptop_form import BR, build_mapping

    two_disks = dict(SAMPLE_DATA)
    two_disks["disks"] = [
        {"name": "SSD", "size_gb": 512, "type": "NVMe"},
        {"name": "HDD", "size_gb": 1024, "type": "SATA"},
    ]
    m = build_mapping(two_disks, None)
    assert m["hostname"] == "LT-0042"
    assert m["serial_number"] == "SN-XYZ-9911"
    assert BR in m["disks"]
    assert "SSD — 512 GB — NVMe" in m["disks"]
    assert "Google Chrome (installed, v131.0.6778.85)" in m["software_checks"]
    assert "Licensed" == m["windows_activation_status"]
    assert "XYZ12" == m["windows_activation_partial_key"]

    single = build_mapping(SAMPLE_DATA, None)     # one disk → no separator
    assert BR not in single["disks"]
    assert single["disks"] == "Samsung SSD — 512 GB — NVMe"


def test_build_mapping_short_specs():
    """Short-form tokens for the corporate form (i7 style CPU, integer RAM)."""
    from app.services.laptop_form import build_mapping
    data = dict(SAMPLE_DATA, cpu="13th Gen Intel(R) Core(TM) i7-1355U",
                ram_gb=16.0)
    m = build_mapping(data, None)
    assert m["cpu_family"] == "i7"
    assert m["ram"] == "16"
    assert m["os_full"] == "Windows 11 Pro 10.0.26100 (64-bit)"
    assert m["disk1"].startswith("Samsung SSD")

    for cpu, family in [
        ("Intel(R) Core(TM) Ultra 7 155H", "Ultra 7"),
        ("AMD Ryzen 5 PRO 8640U", "Ryzen 5"),
        ("Intel(R) Xeon(R) Gold 6338", "Xeon"),
        ("Apple M1 Pro", "Apple M1"),
        ("Intel(R) Core(TM) i5-1135G7", "i5"),
        ("12th Gen Intel(R) Core(TM) i9-12900H", "i9"),
    ]:
        assert build_mapping({"cpu": cpu}, None)["cpu_family"] == family, cpu
    assert build_mapping({"ram_gb": "8"}, None)["ram"] == "8"


def test_fill_document_blanks_unknown_tokens(tmp_path):
    import docx
    from app.services.laptop_form import fill_document
    doc = docx.Document()
    doc.add_paragraph("Value: {{totally_bogus_field}}")
    template = tmp_path / "t.docx"
    doc.save(str(template))
    out = tmp_path / "out.docx"
    result = fill_document(SAMPLE_DATA, None, str(template), str(out))
    assert result["tokens_missing"] == ["totally_bogus_field"]
    back = docx.Document(str(out))
    assert "{{" not in back.paragraphs[0].text      # blanked, not printed raw


def test_fill_document_tokens_split_across_runs(tmp_path):
    import docx
    from app.services.laptop_form import fill_document
    doc = docx.Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Host: {{hos")       # Word sometimes splits tokens
    paragraph.add_run("tname}} in two runs")
    template = tmp_path / "t.docx"
    doc.save(str(template))
    out = tmp_path / "out.docx"
    result = fill_document(SAMPLE_DATA, None, str(template), str(out))
    assert result["tokens_replaced"] >= 1
    back = docx.Document(str(out))
    assert back.paragraphs[0].text == "Host: LT-0042 in two runs"


def test_fill_document_shipped_template_fully_consumed(tmp_path):
    """The starter template (scripts/make_laptop_form_template.py) has no
    token that build_mapping does not supply, and all values land."""
    import docx
    from app.services.laptop_form import fill_document
    out = tmp_path / "out.docx"
    result = fill_document(SAMPLE_DATA, None, "templates/Laptop Form.docx",
                           str(out))
    assert result["tokens_missing"] == []
    assert result["tokens_replaced"] >= 25
    document = docx.Document(str(out))
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    joined = "\n".join(texts)
    assert "{{" not in joined
    for expected in ("LT-0042", "SN-XYZ-9911", "Licensed", "Google Chrome"):
        assert expected in joined


def test_autofill_corporate_form_label_tables(tmp_path):
    """The real B.TECH Finance form has NO {{tokens}} — plain label cells:

        Model            | (empty)          →  model value
        Serial           | (empty)          →  serial value
        NO. | Processor  | (empty) | CONFIG →  cpu short family
        NO. | RAM        | (empty) | CONFIG →  integer RAM
    """
    from docx import Document
    from app.services.laptop_form import fill_document

    doc = Document()
    upper = doc.add_table(rows=2, cols=2)
    upper.style = "Table Grid"
    upper.cell(0, 0).paragraphs[0].add_run("Model")
    upper.cell(1, 0).paragraphs[0].add_run("Serial")

    spec = doc.add_table(rows=3, cols=4)
    spec.style = "Table Grid"
    for i, h in enumerate(["NO.", "SPECIFICATION", "BRAND / TYPE", "CONFIG."]):
        spec.cell(0, i).paragraphs[0].add_run(h).bold = True
    spec.cell(1, 0).paragraphs[0].add_run("01")
    spec.cell(1, 1).paragraphs[0].add_run("Processor")
    spec.cell(2, 0).paragraphs[0].add_run("02")
    spec.cell(2, 1).paragraphs[0].add_run("RAM")

    # second table whose HEADER cell is labeled 'Serial' — must NOT be filled
    second = doc.add_table(rows=2, cols=4)
    second.style = "Table Grid"
    for i, h in enumerate(["NO.", "SPECIFICATION", "BRAND / TYPE", "Serial"]):
        second.cell(0, i).paragraphs[0].add_run(h).bold = True

    # an alternative table whose Serial value cell is PRE-FILLED — untouched
    third = doc.add_table(rows=1, cols=2)
    third.style = "Table Grid"
    third.cell(0, 0).paragraphs[0].add_run("Serial ")
    third.cell(0, 1).paragraphs[0].add_run("TYPED-IN-ADVANCE")

    template = tmp_path / "corporate.docx"
    doc.save(str(template))
    out = tmp_path / "filled.docx"

    data = dict(SAMPLE_DATA, cpu="Intel(R) Core(TM) Ultra 7 155H", ram_gb=32.0)
    result = fill_document(data, None, str(template), str(out))
    assert result["tokens_replaced"] == 0            # template has no tokens
    assert result["labels_filled"] == 4              # Model+Serial+CPU+RAM
    assert result["tokens_missing"] == []

    back = Document(str(out))
    assert back.tables[0].cell(0, 1).text.strip() == "Latitude 5440"
    assert back.tables[0].cell(1, 1).text.strip() == "SN-XYZ-9911"
    assert back.tables[1].cell(1, 2).text.strip() == "Ultra 7"
    assert back.tables[1].cell(2, 2).text.strip() == "32"
    # untouched template content
    assert back.tables[1].cell(1, 0).text.strip() == "01"          # numbering
    assert back.tables[1].cell(0, 2).text.strip() == "BRAND / TYPE"
    assert back.tables[2].cell(0, 3).text.strip() == "Serial"      # header intact
    assert back.tables[2].cell(1, 3).text.strip() == ""            # rows blank
    assert back.tables[3].cell(0, 1).text.strip() == "TYPED-IN-ADVANCE"


def test_autofill_disabled_by_config(tmp_path, monkeypatch):
    import docx as _docx
    from docx import Document
    from app.services.laptop_form import fill_document
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).paragraphs[0].add_run("Serial")
    template = tmp_path / "t.docx"
    doc.save(str(template))
    out = tmp_path / "o.docx"
    import app.config as config
    monkeypatch.setattr(config, "PRINT_AUTOFILL_LABELS", False)
    result = fill_document(SAMPLE_DATA, None, str(template), str(out))
    assert result["labels_filled"] == 0
    assert Document(str(out)).tables[0].cell(0, 1).text.strip() == ""


def test_spawned_docx_multiline_breaks(tmp_path):
    """Multi-line values become real Word line breaks (<w:br/>), not '\n'."""
    from docx import Document
    from docx.oxml.ns import qn
    from app.services.laptop_form import fill_document

    doc = Document()
    doc.add_table(rows=1, cols=1)
    doc.tables[0].style = "Table Grid"
    doc.tables[0].rows[0].cells[0].paragraphs[0].add_run("{{disks}}")
    template = tmp_path / "t.docx"
    doc.save(str(template))

    out = tmp_path / "out.docx"
    fill_document(SAMPLE_DATA, None, str(template), str(out))       # 1 disk
    run = Document(str(out)).tables[0].rows[0].cells[0].paragraphs[0].runs[0]
    assert run.text == "Samsung SSD — 512 GB — NVMe"
    assert len(run._element.findall(qn("w:br"))) == 0

    sample_two_disks = dict(SAMPLE_DATA)
    sample_two_disks["disks"] = [
        {"name": "SSD", "size_gb": 512, "type": "NVMe"},
        {"name": "HDD", "size_gb": 1024, "type": "SATA"},
    ]
    fill_document(sample_two_disks, None, str(template), str(out))  # 2 disks
    run = Document(str(out)).tables[0].rows[0].cells[0].paragraphs[0].runs[0]
    texts = run._element.findall(qn("w:t"))
    assert texts[0].text == "SSD — 512 GB — NVMe"                   # first line
    assert texts[-1].text == "HDD — 1024 GB — SATA"                 # second line
    # exactly one <w:br/> between the two lines (rendered as a line break)
    breaks = run._element.findall(qn("w:br"))
    assert len(breaks) == 1

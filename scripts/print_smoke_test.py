# Smoke test for the Print Laptop Form pipeline (dev machine, no admin needed).
#德 注意: writes only to var/print/_smoke*, deletes nothing else.
import json
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_here))   # repo root (scripts/..)

print("== 1. import & app factory ==")
from app.main import create_app  # noqa: E402
app = create_app()
from app.routers import printing as _pr
print("print routes:", [getattr(r, "path", None) for r in _pr.router.routes])

print("== 2. real printer enumeration (Get-Printer) ==")
from app.services import printing
printers = printing.enumerate_printers()
for p in printers:
    print("  queue:", p)
assert any("HP LaserJet" in (p["name"] or "") for p in printers), "expected the HP queue"

print("== 3. real template fill (python-docx -> docx back-read) ==")
from app.services import laptop_form

sample = {
    "hostname": "LT-0042", "username": "j.doe",
    "serial": "SN-XYZ-9911", "manufacturer": "Dell Inc.", "model": "Latitude 5440",
    "cpu": "Intel i7-1365U", "cpu_cores": 10, "ram_gb": 16,
    "disks": [{"name": "Samsung SSD", "size_gb": 512, "type": "NVMe"}],
    "os_name": "Windows 11 Pro", "os_version": "10.0.26100", "os_arch": "64-bit",
    "ip_addresses": [{"iface": "Ethernet", "ip": "10.42.1.7"}],
    "agent_os": "windows", "collected_at": "2026-10-20T10:00:00+00:00",
    "software_checks": [
        {"id": "chrome", "name": "Google Chrome", "installed": True,
         "version": "131.0.6778.85", "method": "registry"},
        {"id": "vlc", "name": "VLC", "installed": False, "version": None,
         "method": "registry"},
    ],
    "windows_activation": {"status": "Licensed", "product_name": "Windows 11 Pro",
                           "partial_key": "XYZ12"},
}
os.makedirs("var/print", exist_ok=True)
fill = laptop_form.fill_document(sample, None, "templates/Laptop Form.docx",
                                 "var/print/_smoke.docx")
print("  fill result:", fill)
assert fill["tokens_replaced"] >= 10

import docx
d = docx.Document("var/print/_smoke.docx")
all_text = "\n".join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            all_text += "\n" + cell.text
assert "{{" not in all_text, f"unreplaced tokens: {all_text[all_text.find('{{'):][:80]}"
for expect in ("LT-0042", "SN-XYZ-9911", "Google Chrome (installed, v131.0.6778.85)",
               "VLC (not installed)", "Samsung SSD — 512 GB — NVMe",
               "Licensed", "XYZ12", "10.42.1.7"):
    assert expect in all_text, f"missing expected text: {expect}"
print("  docx content verified (multi-line values rendered as line breaks:",
      "\u2021" not in all_text, ")")

print("== 4. expected actionable failures (tools absent on this dev box) ==")
try:
    laptop_form.convert_to_pdf("var/print/_smoke.docx", "var/print")
    print("  UNEXPECTED: conversion succeeded (LibreOffice present?)")
except laptop_form.LaptopFormError as exc:
    print("  converter error (expected):", str(exc)[:120])
try:
    printing.run_ghostscript("var/print/_smoke.docx", printers[0]["name"])
    print("  UNEXPECTED: ghostscript found")
except printing.PrintingError as exc:
    print("  ghostscript error (expected):", str(exc)[:120])

print("== 5. query_jobs on a real queue (works without GS) ==")
jobs = printing.query_jobs("Microsoft Print to PDF")
print("  jobs on 'Microsoft Print to PDF':", jobs)

print("SMOKE OK")
